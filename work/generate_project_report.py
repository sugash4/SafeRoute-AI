from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path("/Users/sugash/Documents/Codex/2026-10-13/so-x20")
OUTPUTS = ROOT / "outputs"
ASSETS = ROOT / "work" / "report_assets"
DOCX_PATH = OUTPUTS / "SafeRoute_AI_Final_Project_Report.docx"

NAVY = "17365D"
BLUE = "2F75B5"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F4F6F8"
GRID = "D9D9D9"
BLACK = "000000"
WHITE = "FFFFFF"
MINT = "38D9A9"
RED = "E64B5D"
AMBER = "E7A614"

FONT_REGULAR = "/System/Library/Fonts/Supplemental/Arial.ttf"
FONT_BOLD = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"


def pil_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REGULAR, size)


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color: str, width: int = 5) -> None:
    draw.line([start, end], fill=color, width=width)
    x1, y1 = start
    x2, y2 = end
    dx, dy = x2 - x1, y2 - y1
    length = max((dx * dx + dy * dy) ** 0.5, 1)
    ux, uy = dx / length, dy / length
    px, py = -uy, ux
    tip = (x2, y2)
    left = (int(x2 - 16 * ux + 8 * px), int(y2 - 16 * uy + 8 * py))
    right = (int(x2 - 16 * ux - 8 * px), int(y2 - 16 * uy - 8 * py))
    draw.polygon([tip, left, right], fill=color)


def make_architecture(path: Path) -> None:
    image = Image.new("RGB", (1600, 720), "white")
    draw = ImageDraw.Draw(image)
    title = pil_font(40, True)
    label = pil_font(28, True)
    detail = pil_font(21)
    draw.text((70, 45), "SafeRoute AI system architecture", font=title, fill="#17365D")
    boxes = [
        (70, 220, 360, 470, "User controls", ["Scenario", "Start and exit", "Penalty weights"]),
        (465, 220, 755, 470, "Campus graph", ["Nodes and edges", "Hazard values", "Blocked corridors"]),
        (860, 220, 1150, 470, "AI planning", ["BFS and DFS", "Greedy and A star", "Replanning and agents"]),
        (1255, 220, 1545, 470, "Explainable result", ["Selected route", "Edge costs", "Metrics and proof"]),
    ]
    for x0, y0, x1, y1, heading, lines in boxes:
        draw.rounded_rectangle((x0, y0, x1, y1), radius=20, fill="#EAF2F8", outline="#2F75B5", width=4)
        draw.text((x0 + 28, y0 + 38), heading, font=label, fill="#17365D")
        for index, line in enumerate(lines):
            draw.text((x0 + 30, y0 + 105 + index * 45), f"• {line}", font=detail, fill="#263238")
    for first, second in zip(boxes, boxes[1:]):
        arrow(draw, (first[2] + 12, 345), (second[0] - 12, 345), "#38A98A", 8)
    draw.text((70, 565), "Observe  →  update graph  →  search or replan  →  explain the selected path", font=pil_font(28, True), fill="#2F75B5")
    image.save(path, quality=95)


POSITIONS = {
    "North Gate": (120, 140), "Library": (390, 140), "Lecture Hall": (660, 140), "East Exit": (1020, 140),
    "Cafeteria": (120, 370), "Central Atrium": (390, 370), "Computer Lab": (660, 370), "Admin Block": (1020, 370),
    "Sports Hall": (120, 600), "South Junction": (390, 600), "Science Lab": (660, 600), "South Exit": (1020, 600),
}

EDGES = [
    ("North Gate", "Library"), ("Library", "Lecture Hall"), ("Lecture Hall", "East Exit"),
    ("North Gate", "Cafeteria"), ("Library", "Central Atrium"), ("Lecture Hall", "Computer Lab"),
    ("East Exit", "Admin Block"), ("Cafeteria", "Central Atrium"), ("Central Atrium", "Computer Lab"),
    ("Computer Lab", "Admin Block"), ("Cafeteria", "Sports Hall"), ("Central Atrium", "South Junction"),
    ("Computer Lab", "Science Lab"), ("Admin Block", "South Exit"), ("Sports Hall", "South Junction"),
    ("South Junction", "Science Lab"), ("Science Lab", "South Exit"), ("Library", "Computer Lab"),
    ("Central Atrium", "Science Lab"),
]


def label_box(draw: ImageDraw.ImageDraw, center: tuple[int, int], text: str) -> None:
    font = pil_font(20, True)
    bbox = draw.textbbox((0, 0), text, font=font)
    width = bbox[2] - bbox[0] + 20
    height = bbox[3] - bbox[1] + 14
    x, y = center
    draw.rounded_rectangle((x - width // 2, y - height // 2, x + width // 2, y + height // 2), radius=8, fill="#D9F7EC", outline="#17365D", width=2)
    draw.text((x - width // 2 + 10, y - height // 2 + 5), text, font=font, fill="#17365D")


def make_route_map(path: Path) -> None:
    image = Image.new("RGB", (1160, 780), "white")
    draw = ImageDraw.Draw(image)
    draw.text((40, 25), "Selected route with node hazards and weighted edge costs", font=pil_font(32, True), fill="#17365D")
    for source, target in EDGES:
        draw.line([POSITIONS[source], POSITIONS[target]], fill="#AAB7C4", width=3)
    route = ["Library", "Computer Lab", "Science Lab", "South Exit"]
    costs = [7.50, 5.20, 4.80]
    offsets = [(0, -28), (55, 0), (0, -28)]
    for index, (source, target) in enumerate(zip(route, route[1:])):
        p0, p1 = POSITIONS[source], POSITIONS[target]
        draw.line([p0, p1], fill="#17A673", width=10)
        midpoint = ((p0[0] + p1[0]) // 2 + offsets[index][0], (p0[1] + p1[1]) // 2 + offsets[index][1])
        label_box(draw, midpoint, f"cost {costs[index]:.2f}")
    for name, (x, y) in POSITIONS.items():
        selected = name in route
        fill = "#2F75B5" if selected else "#B9C5D0"
        draw.ellipse((x - 14, y - 14, x + 14, y + 14), fill=fill, outline="#17365D", width=2)
        label = f"{name}\nhazard 0%"
        draw.multiline_text((x, y - 63), label, font=pil_font(16, selected), fill="#17212B", anchor="mm", align="center", spacing=2)
    draw.text((40, 735), "Route total: 7.50 + 5.20 + 4.80 = 17.50", font=pil_font(24, True), fill="#17365D")
    image.save(path, quality=95)


def make_algorithm_chart(path: Path) -> None:
    image = Image.new("RGB", (1400, 760), "white")
    draw = ImageDraw.Draw(image)
    draw.text((65, 38), "Normal operation algorithm comparison", font=pil_font(38, True), fill="#17365D")
    algorithms = ["BFS", "DFS", "Greedy", "A star"]
    costs = [18.8, 18.8, 18.7, 17.5]
    nodes = [12, 6, 4, 7]
    colors = ["#91A8BE", "#91A8BE", "#5B9BD5", "#38A98A"]
    origin_x, origin_y = 130, 620
    max_height = 420
    scale = max_height / 20
    draw.line((origin_x, 170, origin_x, origin_y), fill="#4C5B67", width=3)
    draw.line((origin_x, origin_y, 1320, origin_y), fill="#4C5B67", width=3)
    for value in range(0, 21, 5):
        y = origin_y - value * scale
        draw.line((origin_x - 8, y, 1320, y), fill="#E1E6EA", width=2)
        draw.text((60, y - 14), str(value), font=pil_font(20), fill="#4C5B67")
    bar_width = 180
    gap = 90
    for index, (name, cost, expanded, color) in enumerate(zip(algorithms, costs, nodes, colors)):
        x0 = origin_x + 100 + index * (bar_width + gap)
        x1 = x0 + bar_width
        y0 = origin_y - cost * scale
        draw.rounded_rectangle((x0, y0, x1, origin_y), radius=12, fill=color)
        draw.text(((x0 + x1) // 2, y0 - 34), f"{cost:.1f}", font=pil_font(25, True), fill="#17365D", anchor="mm")
        draw.text(((x0 + x1) // 2, origin_y + 38), name, font=pil_font(24, True), fill="#17212B", anchor="mm")
        draw.text(((x0 + x1) // 2, origin_y + 75), f"{expanded} nodes expanded", font=pil_font(18), fill="#4C5B67", anchor="mm")
    draw.text((65, 705), "Weighted cost (lower is better); default penalties: hazard 5.0 and congestion 2.0", font=pil_font(21), fill="#4C5B67")
    image.save(path, quality=95)


def make_replanning(path: Path) -> None:
    image = Image.new("RGB", (1500, 620), "white")
    draw = ImageDraw.Draw(image)
    draw.text((60, 35), "Dynamic replanning after an unexpected corridor closure", font=pil_font(38, True), fill="#17365D")
    steps = [
        (60, 190, 340, 410, "1  Initial plan", ["Library", "Computer Lab", "Science Lab", "South Exit"]),
        (460, 190, 740, 410, "2  New observation", ["Agent at Computer Lab", "Science Lab corridor", "becomes blocked"]),
        (860, 190, 1140, 410, "3  Graph update", ["Remove blocked action", "Keep current state", "Run A star again"]),
        (1260, 190, 1440, 410, "4  New plan", ["Computer Lab", "Admin Block", "South Exit"]),
    ]
    for x0, y0, x1, y1, heading, lines in steps:
        draw.rounded_rectangle((x0, y0, x1, y1), radius=18, fill="#F4F6F8", outline="#2F75B5", width=4)
        draw.text((x0 + 22, y0 + 25), heading, font=pil_font(24, True), fill="#17365D")
        for i, line in enumerate(lines):
            color = "#C23B4B" if "blocked" in line else "#27323B"
            draw.text((x0 + 22, y0 + 85 + i * 36), line, font=pil_font(19), fill=color)
    for first, second in zip(steps, steps[1:]):
        arrow(draw, (first[2] + 15, 300), (second[0] - 15, 300), "#38A98A", 7)
    draw.text((60, 500), "The new search starts at Computer Lab because the agent has already moved from the Library.", font=pil_font(25, True), fill="#17365D")
    image.save(path, quality=95)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 100, start: int = 110, bottom: int = 100, end: int = 110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), GRID)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def set_run_font(run, name: str = "Arial", size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_body(doc: Document, text: str, bold_lead: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    paragraph.style = doc.styles["Normal"]
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        set_run_font(paragraph.add_run(item))


def add_numbered(doc: Document, items: list[str]) -> None:
    for index, item in enumerate(items, start=1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.65)
        paragraph.paragraph_format.first_line_indent = Cm(-0.65)
        set_run_font(paragraph.add_run(f"{index}.  {item}"))


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True


def add_formula(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(7)
    paragraph.paragraph_format.space_after = Pt(7)
    run = paragraph.add_run(text)
    set_run_font(run, name="Arial", size=11.5, bold=True, color=NAVY)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    repeat_table_header(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(paragraph.add_run(value), size=9, bold=True, color=WHITE)
        if widths:
            cell.width = Inches(widths[index])
    for row_index, values in enumerate(rows):
        row = table.add_row()
        if row_index % 2:
            for cell in row.cells:
                set_cell_shading(cell, PALE_BLUE)
        for index, value in enumerate(values):
            cell = row.cells[index]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(value) < 18 else WD_ALIGN_PARAGRAPH.LEFT
            set_run_font(paragraph.add_run(value), size=8.5)
            if widths:
                cell.width = Inches(widths[index])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_figure(doc: Document, path: Path, caption: str, width: float = 6.7) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    picture = run.add_picture(str(path), width=Inches(width))
    # Store a useful description in the DOCX drawing metadata so screen readers
    # do not encounter an unexplained image.
    picture._inline.docPr.set("descr", caption)
    picture._inline.docPr.set("title", caption)
    caption_paragraph = doc.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.keep_with_next = False
    caption_run = caption_paragraph.add_run(caption)
    set_run_font(caption_run, size=9, bold=True, color="4C5B67")


def page_break(doc: Document) -> None:
    doc.add_page_break()


def build_document() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    OUTPUTS.mkdir(parents=True, exist_ok=True)
    architecture = ASSETS / "architecture.png"
    route_map = ASSETS / "route_map.png"
    comparison = ASSETS / "algorithm_comparison.png"
    replanning = ASSETS / "replanning.png"
    make_architecture(architecture)
    make_route_map(route_map)
    make_algorithm_chart(comparison)
    make_replanning(replanning)

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("202830")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for style_name, size, before, after in (("Title", 28, 0, 12), ("Heading 1", 18, 12, 7), ("Heading 2", 13, 9, 5), ("Heading 3", 11, 7, 4)):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    add_page_number(footer_paragraph)
    set_run_font(footer_paragraph.runs[0], size=8, color="6B747C")

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(75)
    set_run_font(title.add_run("SafeRoute AI Intelligent Campus Evacuation Planner"), size=28, bold=True, color=BLACK)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(42)
    set_run_font(subtitle.add_run("Academic Project Report"), size=16, bold=True, color=NAVY)
    course = doc.add_paragraph()
    course.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(course.add_run("Course  Artificial Intelligence"), size=13, bold=True)
    doc.add_paragraph()
    cover_fields = [
        ["Student name", "Sugash K"],
        ["Register number", "24BCE2565"],
    ]
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for label, blank in cover_fields:
        row = table.add_row()
        row.cells[0].width = Inches(1.8)
        row.cells[1].width = Inches(3.5)
        for cell in row.cells:
            set_cell_margins(cell, top=90, bottom=90)
        set_run_font(row.cells[0].paragraphs[0].add_run(label), size=10.5, bold=True, color=NAVY)
        set_run_font(row.cells[1].paragraphs[0].add_run(blank), size=10.5)
    closing = doc.add_paragraph()
    closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
    closing.paragraph_format.space_before = Pt(55)
    set_run_font(closing.add_run("Submitted in partial fulfilment of the academic course requirements"), size=10, color="4C5B67")

    page_break(doc)
    add_heading(doc, "Abstract")
    add_body(doc, "SafeRoute AI is an explainable campus evacuation planner that applies classical artificial intelligence to route selection in a changing emergency environment. The application represents a campus as a weighted graph and compares Breadth First Search, Depth First Search, Greedy Best First Search, and risk aware A star. Its principal result is that the physically shortest route is not always the best evacuation route: a longer alternative can have a lower weighted cost when hazard probability, congestion, or blocked corridors are considered.")
    add_body(doc, "The system provides five interactive modules: Live map, Compare, Replanning, Multi agent, and Map editor. The Live map displays the hazard value of every node, the weighted cost of each selected route edge, and a mathematical verification showing that the edge costs sum to the reported route cost. Replanning demonstrates how an agent updates its plan after a corridor becomes unavailable. Multi agent planning assigns several evacuees to exits subject to capacity and shared corridor penalties. Nine automated tests verify search correctness, hazard avoidance, blocked edge handling, replanning, map reconstruction, multi agent capacity, probability calculations, and visual cost labels.")
    add_body(doc, "The implementation is an educational prototype rather than a certified life safety system. It demonstrates state space search, informed search, utility based decision making, planning in a dynamic environment, and multi agent coordination in a form that can be inspected and explained during assessment.")
    add_heading(doc, "Keywords", level=2)
    add_body(doc, "Artificial intelligence, state space search, A star, graph search, evacuation planning, dynamic replanning, multi agent planning, hazard aware routing, explainable AI")
    add_heading(doc, "Report structure", level=2)
    add_bullets(doc, [
        "Problem definition, objectives, scope, and course concept mapping",
        "System architecture, data representation, and mathematical cost model",
        "Search algorithms and functional modules",
        "Experimental results, automated testing, limitations, and future work",
        "Appendices containing the project structure and local execution procedure",
    ])

    page_break(doc)
    add_heading(doc, "1 Introduction")
    add_heading(doc, "1 1 Background", level=2)
    add_body(doc, "Evacuation route selection is commonly described as a shortest path problem, but emergency movement involves more than distance. Smoke, chemical leaks, crowd density, and corridor closures can make a short path unsafe or unusable. A useful planner must represent the environment, compare alternative actions, incorporate changing costs, and produce a new plan when its assumptions become invalid.")
    add_heading(doc, "1 2 Problem statement", level=2)
    add_body(doc, "The project addresses the following question: how can a goal based agent choose and explain a campus evacuation route when distance, hazard probability, congestion, and unexpected corridor closures affect the quality of each action?")
    add_heading(doc, "1 3 Objectives", level=2)
    add_numbered(doc, [
        "Represent campus locations and corridors as a searchable weighted graph.",
        "Implement BFS, DFS, Greedy Best First Search, and A star directly in Python.",
        "Combine distance, average node hazard, and corridor congestion in one transparent cost function.",
        "Show route edge costs and node hazard values directly on the Live map.",
        "Demonstrate replanning after an unexpected environmental change.",
        "Coordinate multiple agents while respecting exit capacity and route conflicts.",
        "Allow a user to edit and validate a custom campus graph.",
        "Verify important behaviour with reproducible experiments and automated tests.",
    ])
    add_heading(doc, "1 4 Scope", level=2)
    add_body(doc, "The scope is limited to a small simulated campus graph with manually defined hazard and congestion values. The planner calculates recommendations but does not connect to physical sensors, control doors, or issue real emergency instructions. The interface is designed for academic demonstration and algorithm comparison.")

    page_break(doc)
    add_heading(doc, "2 Course Concept Mapping")
    add_table(doc,
        ["Course concept", "Project implementation", "Evidence in application"],
        [
            ["State space representation", "CampusGraph stores nodes, edges, hazards, congestion, and blocked corridors.", "Live map and Map editor"],
            ["Uninformed search", "BFS uses a queue and DFS uses a stack.", "Compare page"],
            ["Informed search", "Greedy uses h(n); A star uses g(n) + h(n).", "Search strategy control and Compare page"],
            ["Intelligent agents", "A goal based, utility sensitive agent selects an exit route.", "Route result and weighted cost"],
            ["Planning", "The plan is recomputed from the current location after a corridor closes.", "Replanning page"],
            ["Multi agent planning", "Joint exit assignment includes capacity and shared corridor penalties.", "Multi agent page"],
            ["Knowledge representation", "Hazard probabilities and corridor properties encode environmental state.", "Scenario presets and Map editor"],
            ["Evaluation", "Algorithms are compared by cost, distance, risk, congestion, expanded nodes, and runtime.", "Compare page and experiment output"],
        ],
        widths=[1.55, 3.1, 1.7],
    )
    add_heading(doc, "2 1 Agent classification", level=2)
    add_body(doc, "SafeRoute AI is a goal based and utility sensitive agent. Reaching the selected exit is the goal; minimizing the weighted route cost is the utility objective. The environment is sequential because one movement changes the next state, dynamic because conditions can change during execution, partially observable because future changes are not known in advance, and nondeterministic because a corridor may unexpectedly become unavailable.")

    page_break(doc)
    add_heading(doc, "3 System Design")
    add_figure(doc, architecture, "Figure 1  SafeRoute AI processing architecture", width=6.7)
    add_heading(doc, "3 1 Input layer", level=2)
    add_body(doc, "The Streamlit interface collects the emergency scenario, starting location, target exit, search strategy, hazard penalty, and congestion penalty. A user can also activate a validated custom map.")
    add_heading(doc, "3 2 Environment model", level=2)
    add_body(doc, "CampusGraph stores coordinates for drawing and heuristic calculation, adjacency lists for connectivity, hazard probabilities for nodes, congestion values for edges, and a set of blocked edges. The neighbours function filters blocked corridors before any search algorithm can expand them.")
    add_heading(doc, "3 3 Planning layer", level=2)
    add_body(doc, "The selected algorithm searches the graph and records parent links, expansion order, accumulated cost, and runtime. A star and the multi agent planner use the weighted cost function during decision making. The replanning module modifies a copy of the graph and searches again from the current location.")
    add_heading(doc, "3 4 Explanation layer", level=2)
    add_body(doc, "Plotly renders the graph, selected path, blocked corridors, node hazards, and route edge costs. Tables and metrics explain the final result. The mathematical proof below the Live map substitutes the current values into the same cost function used by the search implementation.")

    page_break(doc)
    add_heading(doc, "4 Mathematical Model")
    add_heading(doc, "4 1 Graph formulation", level=2)
    add_body(doc, "The campus is an undirected weighted graph G = (V, E). Each vertex v in V is a campus location. Each edge (u, v) in E is a corridor that can be traversed in both directions unless it is blocked. A search state is the current vertex. An action moves to an unblocked neighbouring vertex, and the goal test checks whether the current vertex is the selected exit.")
    add_heading(doc, "4 2 Edge cost", level=2)
    add_formula(doc, "mean_hazard(u,v) = [H(u) + H(v)] / 2")
    add_formula(doc, "C(u,v) = d(u,v) [1 + w_h mean_hazard(u,v) + w_c Q(u,v)]")
    add_body(doc, "Here d is corridor distance, H is node hazard probability, Q is corridor congestion, w_h is the user selected hazard penalty, and w_c is the congestion penalty. All terms are nonnegative, so the weighted cost is never lower than the physical distance.")
    add_heading(doc, "4 3 Default route proof", level=2)
    add_table(doc,
        ["Selected edge", "Substitution", "Cost"],
        [
            ["Library to Computer Lab", "5.00 [1 + 5.0(0.00) + 2.0(0.25)]", "7.50"],
            ["Computer Lab to Science Lab", "4.00 [1 + 5.0(0.00) + 2.0(0.15)]", "5.20"],
            ["Science Lab to South Exit", "4.00 [1 + 5.0(0.00) + 2.0(0.10)]", "4.80"],
        ],
        widths=[2.15, 3.55, 0.7],
    )
    add_formula(doc, "Route cost = 7.50 + 5.20 + 4.80 = 17.50")
    add_body(doc, "The Live map calculates these values from the active graph rather than displaying fixed labels. Changing the scenario or either penalty causes Streamlit to rerun the search, update the path, and regenerate the proof.")
    add_heading(doc, "4 4 A star evaluation function", level=2)
    add_formula(doc, "f(n) = g(n) + h(n)")
    add_body(doc, "The value g(n) is the best known accumulated weighted cost from the start to n. The heuristic h(n) is straight line distance from n to the exit. Since weighted edge cost is at least physical distance, the straight line estimate remains a lower bound when edge distances are consistent with the map. Under this assumption the heuristic is admissible.")

    page_break(doc)
    add_heading(doc, "5 Search Algorithms")
    add_table(doc,
        ["Algorithm", "Frontier", "Selection rule", "Main limitation"],
        [
            ["BFS", "FIFO queue", "Shallowest unvisited state", "Minimizes edges, not weighted safety cost"],
            ["DFS", "LIFO stack", "Deepest available branch", "Result depends on neighbour order"],
            ["Greedy", "Priority queue", "Lowest h(n)", "Ignores accumulated cost"],
            ["A star", "Priority queue", "Lowest g(n) + h(n)", "Quality depends on graph and heuristic assumptions"],
        ],
        widths=[1.0, 1.25, 2.1, 2.05],
    )
    add_heading(doc, "5 1 Breadth First Search", level=2)
    add_body(doc, "BFS expands the shallowest state first and is complete on this finite graph. With a visited set and adjacency list its graph search complexity is O(V + E). It is optimal only for the fewest edge objective when all actions have equal cost.")
    add_heading(doc, "5 2 Depth First Search", level=2)
    add_body(doc, "DFS follows one branch before backtracking. The iterative implementation uses a Python list as a stack. It can reach the exit with a small frontier, but the first solution need not be short, safe, or low cost.")
    add_heading(doc, "5 3 Greedy Best First Search", level=2)
    add_body(doc, "Greedy search prioritizes only the estimated remaining distance. It often expands few nodes but can enter a route with high accumulated risk because it does not include g(n) in its priority.")
    add_heading(doc, "5 4 Risk aware A star", level=2)
    add_body(doc, "A star maintains a g score for each discovered node, updates a parent map when a cheaper path is found, and orders its heap by g(n) + h(n). This is the primary risk aware strategy because its accumulated cost includes distance, mean hazard, and congestion.")

    page_break(doc)
    add_heading(doc, "6 Functional Modules")
    add_figure(doc, route_map, "Figure 2  Live map with hazard labels and selected edge costs", width=6.7)
    add_heading(doc, "6 1 Live map", level=2)
    add_body(doc, "The Live map shows the selected route as a mint line. Every node label displays its hazard probability, and each selected edge has a direct weighted cost label. Hover information provides distance, congestion, average hazard, and weighted cost. The metrics below the map report total weighted cost, physical distance, nodes expanded, and runtime.")
    add_heading(doc, "6 2 Compare", level=2)
    add_body(doc, "The Compare page runs all four algorithms on the same graph and presents reached status, weighted cost, distance, risk exposure, congestion exposure, expanded nodes, runtime, and route. This separates route quality from search effort and prevents a single metric from being treated as universally decisive.")
    add_heading(doc, "6 3 Map editor", level=2)
    add_body(doc, "The Map editor exposes node and corridor tables. Validation rejects empty or duplicate node names, missing edge endpoints, nonpositive distances, and probability values outside zero to one. The custom graph remains separate from the reproducible scenario presets until the user activates it.")

    page_break(doc)
    add_heading(doc, "7 Dynamic Replanning and Multi Agent Planning")
    add_figure(doc, replanning, "Figure 3  Observe update and replan sequence", width=6.7)
    add_heading(doc, "7 1 Replanning procedure", level=2)
    add_numbered(doc, [
        "Calculate the initial route from the selected start to the exit.",
        "Simulate movement to the next location on that route.",
        "Observe that a corridor in the remaining plan has become blocked.",
        "Add the corridor to the blocked edge set.",
        "Discard the invalid remainder and run A star from the current location.",
        "Display the revised route or report that no alternative remains.",
    ])
    add_heading(doc, "7 2 Multi agent objective", level=2)
    add_formula(doc, "J = sum(route_cost_i) + conflict_penalty x shared_corridor_uses")
    add_body(doc, "The centralized planner evaluates feasible assignments of selected agents to available exits. It enforces a maximum number of agents per exit and increases the joint objective when routes reuse the same corridor. For the default three agent experiment, the objective is 13.60 with zero shared corridors.")
    add_table(doc,
        ["Agent", "Start", "Assigned exit", "Cost", "Route"],
        [
            ["Agent 1", "Library", "North Gate", "3.60", "Library to North Gate"],
            ["Agent 2", "Cafeteria", "North Gate", "5.20", "Cafeteria to North Gate"],
            ["Agent 3", "Science Lab", "South Exit", "4.80", "Science Lab to South Exit"],
        ],
        widths=[0.8, 1.1, 1.2, 0.65, 2.65],
    )

    page_break(doc)
    add_heading(doc, "8 Implementation")
    add_heading(doc, "8 1 Technology stack", level=2)
    add_table(doc,
        ["Technology", "Purpose"],
        [
            ["Python", "Core implementation language"],
            ["Streamlit", "Interactive local web interface"],
            ["Plotly", "Campus graph and route visualization"],
            ["Pandas", "Algorithm comparison and editable tables"],
            ["Pytest", "Automated correctness verification"],
        ],
        widths=[1.6, 4.8],
    )
    add_heading(doc, "8 2 Important data structures", level=2)
    add_bullets(doc, [
        "Adjacency dictionary for corridor connectivity",
        "Deque for the BFS frontier",
        "Python list for the DFS stack",
        "Heap based priority queue for Greedy and A star",
        "Parent dictionary for route reconstruction",
        "g score dictionary for the cheapest accumulated A star cost",
        "Set of unordered node pairs for blocked corridors",
    ])
    add_heading(doc, "8 3 Explainability", level=2)
    add_body(doc, "The project intentionally avoids using NetworkX to calculate routes. Direct implementations make the queue, stack, heap, parent map, heuristic, and cost updates visible in the source code. The Live map and proof section connect those internal calculations to values that can be inspected during presentation.")
    add_heading(doc, "8 4 Project organization", level=2)
    add_body(doc, "The user interface is contained in app.py. The evacuation package separates models, algorithms, scenarios, planning, multi agent logic, risk utilities, and visualization. Scripts generate reproducible experiment results, while the tests directory contains automated verification.")

    page_break(doc)
    add_heading(doc, "9 Results and Discussion")
    add_figure(doc, comparison, "Figure 4  Default scenario weighted cost and search effort", width=6.7)
    add_table(doc,
        ["Algorithm", "Weighted cost", "Distance", "Nodes expanded", "Selected route summary"],
        [
            ["BFS", "18.80", "13.0", "12", "Via Central Atrium and Science Lab"],
            ["DFS", "18.80", "13.0", "6", "Via Central Atrium and Science Lab"],
            ["Greedy", "18.70", "13.0", "4", "Via Computer Lab and Admin Block"],
            ["A star", "17.50", "13.0", "7", "Via Computer Lab and Science Lab"],
        ],
        widths=[0.85, 1.05, 0.8, 1.0, 2.7],
    )
    add_body(doc, "Under normal operation, all reported routes have the same physical distance, but their congestion exposure differs. A star returns the lowest weighted cost of 17.50. Greedy expands fewer nodes, which demonstrates that search effort and route quality are different evaluation dimensions.")
    add_heading(doc, "9 1 Hazard scenario evidence", level=2)
    add_body(doc, "During the Science lab leak scenario, the Science Lab hazard is 0.98, South Junction is 0.45, and Computer Lab is 0.30. With the standard penalties, BFS and DFS continue through the hazardous southern region and return a weighted cost of 40.85. A star instead selects Library to Lecture Hall to East Exit to Admin Block to South Exit, a longer 15.0 distance route with a lower weighted cost of 19.10 and zero modeled risk exposure.")
    add_heading(doc, "9 2 Interpretation", level=2)
    add_body(doc, "The results support the project objective: minimizing geometric distance or expanded nodes alone does not guarantee the best safety aware route. The weighted objective allows the agent to trade additional distance for reduced hazard and congestion while remaining explainable through per edge calculations.")

    add_heading(doc, "10 Testing and Verification")
    add_body(doc, "The final implementation passes nine automated tests. The tests are deterministic and run locally with Pytest.")
    add_table(doc,
        ["Test area", "Expected behaviour"],
        [
            ["All search algorithms", "Every algorithm reaches the exit on the base graph."],
            ["Severe Library fire", "A star avoids the high hazard Library when a safer alternative exists."],
            ["Blocked corridors", "Returned paths never traverse an edge marked as blocked."],
            ["Probability utility", "Positive evidence increases the calculated hazard probability."],
            ["No path case", "The search returns reached false and an empty path."],
            ["Dynamic replanning", "The revised plan excludes the newly blocked corridor."],
            ["Multi agent capacity", "No exit receives more agents than its configured capacity."],
            ["Editable map", "A graph reconstructed from editor tables preserves connectivity."],
            ["Visual calculation labels", "Displayed edge costs sum to the route cost and node hazards are shown."],
        ],
        widths=[1.7, 4.7],
    )
    add_heading(doc, "10 1 Verification results", level=2)
    add_bullets(doc, [
        "Pytest result: 9 passed",
        "Default selected edge labels: 7.50, 5.20, and 4.80",
        "Verified sum: 17.50, equal to the reported weighted route cost",
        "Streamlit health endpoint returns a successful response while the application is running",
        "The packaged project archive passes integrity verification",
    ])

    page_break(doc)
    add_heading(doc, "11 Limitations Ethics and Future Work")
    add_heading(doc, "11 1 Limitations", level=2)
    add_bullets(doc, [
        "The campus graph and scenario values are simulated rather than measured by certified sensors.",
        "Hazard values do not automatically spread over time according to a physical fire or gas model.",
        "The multi agent assignment method enumerates combinations and is suitable only for a small classroom graph.",
        "The model does not yet include mobility restrictions, stairs, smoke visibility, travel speed, or human behaviour.",
        "Runtime measurements on a small graph are illustrative and should not be treated as large scale benchmarks.",
    ])
    add_heading(doc, "11 2 Ethical and safety considerations", level=2)
    add_body(doc, "SafeRoute AI is not a certified life safety product. A real deployment would require validated building maps, calibrated and redundant sensors, accessibility planning, secure communications, privacy controls, failure recovery, human oversight, regulatory approval, and alignment with official emergency procedures. The application must not be used as the sole basis for real evacuation decisions.")
    add_heading(doc, "11 3 Future work", level=2)
    add_bullets(doc, [
        "Integrate real time sensor and crowd density data.",
        "Model time dependent hazard spread and corridor capacity.",
        "Add accessible routing for users with different mobility needs.",
        "Replace exhaustive multi agent assignment with a scalable coordination method.",
        "Provide mobile guidance and offline fallback behaviour.",
        "Evaluate the planner on larger validated building maps.",
    ])
    add_heading(doc, "12 Conclusion")
    add_body(doc, "SafeRoute AI demonstrates how classical artificial intelligence can convert a campus evacuation problem into an explainable state space search task. The project implements four search strategies, a transparent weighted cost model, dynamic replanning, multi agent coordination, editable graph validation, and visible mathematical verification. Experimental results show that A star can choose a safer or less congested route even when another algorithm expands fewer nodes or follows an equally short physical path. The project therefore meets its academic objective while clearly separating an educational simulation from a real safety certified system.")

    page_break(doc)
    add_heading(doc, "References")
    references = [
        "Russell, S. J., and Norvig, P. Artificial Intelligence A Modern Approach. Fourth edition, Pearson, 2021.",
        "Hart, P. E., Nilsson, N. J., and Raphael, B. A Formal Basis for the Heuristic Determination of Minimum Cost Paths. IEEE Transactions on Systems Science and Cybernetics, 1968.",
        "Python Software Foundation. Python 3 Documentation. https://docs.python.org/3/",
        "Streamlit. Streamlit Documentation. https://docs.streamlit.io/",
        "Plotly. Plotly Python Graphing Library Documentation. https://plotly.com/python/",
        "Pandas Development Team. Pandas Documentation. https://pandas.pydata.org/docs/",
        "Pytest Development Team. Pytest Documentation. https://docs.pytest.org/",
    ]
    for index, reference in enumerate(references, start=1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.6)
        paragraph.paragraph_format.first_line_indent = Cm(-0.6)
        set_run_font(paragraph.add_run(f"{index}. {reference}"), size=10)
    add_heading(doc, "Appendix A Project Structure")
    structure = [
        ("app.py", "Streamlit interface and page coordination"),
        ("evacuation/models.py", "Graph, edge, and result data models"),
        ("evacuation/algorithms.py", "BFS, DFS, Greedy, and A star"),
        ("evacuation/scenarios.py", "Campus map and emergency scenarios"),
        ("evacuation/planning.py", "Unexpected blockage replanning"),
        ("evacuation/multiagent.py", "Exit assignment and conflict penalties"),
        ("evacuation/visualization.py", "Plotly map and direct calculation labels"),
        ("tests/", "Automated correctness tests"),
        ("scripts/run_experiments.py", "Reproducible experiment generator"),
    ]
    add_table(doc, ["File or folder", "Responsibility"], [[a, b] for a, b in structure], widths=[2.3, 4.1])
    add_heading(doc, "Appendix B Local Execution")
    add_numbered(doc, [
        "Open a terminal in the project folder.",
        "Create a virtual environment with python3 -m venv .venv.",
        "Activate it with source .venv/bin/activate on macOS or Linux.",
        "Install dependencies with python -m pip install -r requirements.txt.",
        "Run the website with streamlit run app.py.",
        "Open http://localhost:8501 in a browser.",
        "Run tests with pytest -q before submission.",
    ])

    doc.core_properties.title = "SafeRoute AI Intelligent Campus Evacuation Planner"
    doc.core_properties.subject = "Academic Artificial Intelligence project report"
    doc.core_properties.author = "Student"
    doc.core_properties.keywords = "AI, A star, evacuation, graph search, replanning"
    doc.save(DOCX_PATH)


if __name__ == "__main__":
    build_document()
    print(DOCX_PATH)
