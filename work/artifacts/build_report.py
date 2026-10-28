from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "outputs" / "SafeRoute_AI_Project_Report.docx"
SCREENSHOTS = ROOT / "work" / "screenshots"
EQUATIONS = ROOT / "work" / "artifacts" / "equations"

NAVY = "17365D"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F5F7F9"
BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)
ACCENT = RGBColor(23, 54, 93)


def set_cell_fill(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), BORDER)


def repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def prevent_row_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:cantSplit")
    marker.set(qn("w:val"), "true")
    properties.append(marker)


def remove_paragraph_border(paragraph) -> None:
    properties = paragraph._p.get_or_add_pPr()
    border = properties.find(qn("w:pBdr"))
    if border is not None:
        properties.remove(border)


def style_table(table, widths: list[float] | None = None) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    repeat_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        prevent_row_split(row)
        for col_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if widths:
                cell.width = Inches(widths[col_index])
            if row_index == 0:
                set_cell_fill(cell, NAVY)
            elif row_index % 2 == 0:
                set_cell_fill(cell, PALE_BLUE)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.name = "Liberation Sans"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Liberation Sans")
                    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Liberation Sans")
                    run.font.size = Pt(9)
                    if row_index == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)


def add_table(doc, headers: list[str], rows: list[list[object]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = str(header)
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value)
    style_table(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_body(doc, text: str, bold_lead: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    paragraph.style = doc.styles["Body Text"]
    if bold_lead:
        run = paragraph.add_run(bold_lead)
        run.bold = True
    paragraph.add_run(text)


def add_bullets(doc, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_caption(doc, text: str) -> None:
    paragraph = doc.add_paragraph(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.style = doc.styles["Caption"]


def add_equation(doc, filename: str, width: float) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(EQUATIONS / filename), width=Inches(width))


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.footer_distance = Inches(0.35)
    add_page_number(section.footer.paragraphs[0])

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12

    body = styles["Body Text"]
    body.font.name = "Liberation Sans"
    body._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
    body._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
    body.font.size = Pt(11)
    body.font.color.rgb = BLACK
    body.paragraph_format.space_after = Pt(8)
    body.paragraph_format.line_spacing = 1.15

    title = styles["Title"]
    title.font.name = "Liberation Sans"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
    title.font.size = Pt(26)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title_properties = title._element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)

    for name, size, before, after in (("Heading 1", 17, 16, 7), ("Heading 2", 13, 12, 5), ("Heading 3", 11, 9, 3)):
        style = styles[name]
        style.font.name = "Liberation Sans"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    styles["Caption"].font.name = "Liberation Sans"
    styles["Caption"].font.size = Pt(9)
    styles["Caption"].font.italic = True
    styles["Caption"].font.color.rgb = BLACK


def main() -> None:
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    remove_paragraph_border(title)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(90)
    title.add_run("Design and Evaluation of a Risk Aware Intelligent Agent for Dynamic Campus Evacuation")
    subtitle = doc.add_paragraph()
    remove_paragraph_border(subtitle)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(18)
    run = subtitle.add_run("Artificial Intelligence Academic Project Report")
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = ACCENT
    details = doc.add_paragraph()
    details.alignment = WD_ALIGN_PARAGRAPH.CENTER
    details.paragraph_format.space_before = Pt(72)
    details.add_run(
        "Student name  ______________________________\n\n"
        "Register number  ____________________________\n\n"
        "Course and section  __________________________\n\n"
        "Institution  _________________________________\n\n"
        "Faculty guide  _______________________________\n\n"
        "Academic year  ______________________________"
    )
    purpose = doc.add_paragraph()
    purpose.alignment = WD_ALIGN_PARAGRAPH.CENTER
    purpose.paragraph_format.space_before = Pt(48)
    purpose.add_run("Submitted in partial fulfilment of the Artificial Intelligence course requirements")
    doc.add_page_break()

    doc.add_heading("Abstract", level=1)
    add_body(doc, "SafeRoute AI is an educational campus evacuation simulator that demonstrates classical search, probabilistic reasoning, planning under change, multiagent coordination, and genetic optimization in one explainable application. The campus is represented as a weighted graph. The route cost combines distance, hazard probability, and congestion. Four search strategies are compared using route cost, distance, risk exposure, congestion exposure, and nodes expanded. Bayesian updates convert uncertain smoke and heat evidence into hazard probabilities. A replanning drill changes the environment after movement begins. A centralized multiagent planner assigns agents to exits while limiting exit load and penalizing shared corridors. A genetic algorithm selects emergency resource locations by minimizing average risk-aware access cost. The Streamlit application also provides a validated campus-map editor. Twenty experiment runs across five scenarios and nine automated tests support the evaluation. The prototype is designed for academic demonstration and is not a certified emergency system.")
    doc.add_heading("Keywords", level=2)
    add_body(doc, "Artificial intelligence, A star search, Bayesian reasoning, evacuation planning, genetic algorithm, intelligent agent, multiagent planning, state-space search")

    doc.add_heading("Project Scope", level=1)
    add_body(doc, "The project demonstrates how an intelligent agent can choose, explain, revise, and coordinate evacuation routes in a simulated campus. It focuses on transparent algorithms that match the course syllabus. It does not model real smoke propagation, human movement, accessibility needs, or certified emergency procedures.")
    doc.add_heading("Implemented Deliverables", level=2)
    add_bullets(doc, [
        "Manual implementations of BFS, DFS, Greedy Best-First Search, and risk-aware A star",
        "Bayesian smoke and heat sensor update",
        "Dynamic corridor blockage with replanning from the current state",
        "Joint exit assignment for up to five agents",
        "Genetic optimization of emergency resource placement",
        "Validated node and corridor map editor",
        "Interactive visual comparison, reproducible experiments, and automated tests",
    ])

    doc.add_page_break()
    doc.add_heading("Course Outcome Mapping", level=1)
    add_table(doc, ["Course topic", "Project evidence"], [
        ["Intelligent agents and environments", "Goal-based, utility-sensitive agent in a dynamic and partially observable campus"],
        ["State-space search", "Campus graph, successor function, goal test, path reconstruction"],
        ["Uninformed and informed search", "BFS, DFS, Greedy, and A star comparison"],
        ["Uncertain reasoning", "Bayesian update from smoke and heat evidence"],
        ["Planning", "Route execution followed by observation and replanning"],
        ["Multiagent planning", "Joint exit assignment with conflict and capacity constraints"],
        ["Local search", "Genetic algorithm for resource locations"],
        ["Contemporary issues", "Sensor error, model limitations, safety and accountability"],
    ], [2.25, 4.35])
    doc.add_heading("1 Introduction", level=1)
    add_body(doc, "An evacuation route cannot be judged by distance alone. A short corridor may contain smoke, a crowded atrium may delay movement, and an initially valid route may become blocked. SafeRoute AI treats these conditions as parts of an explicit search problem. The application exposes each algorithm's expansion order and route metrics so an evaluator can inspect the decision rather than accept a hidden model output.")
    doc.add_heading("1 1 Problem Statement", level=2)
    add_body(doc, "Given a campus graph, a starting location, one or more exits, uncertain hazard evidence, congestion values, and blocked corridors, determine a feasible route or coordinated set of routes with low weighted evacuation cost. If the environment changes, update the state and calculate a revised plan.")
    doc.add_heading("1 2 Objectives", level=2)
    add_bullets(doc, [
        "Represent a campus as an editable weighted graph",
        "Compare search strategies using observable metrics",
        "Incorporate uncertain sensor evidence into route cost",
        "Respond to an unexpected blocked corridor",
        "Coordinate several agents without overloading one exit",
        "Optimize resource placement through a reproducible genetic search",
    ])

    doc.add_heading("2 Problem Formulation", level=1)
    add_body(doc, "The graph G contains campus locations as vertices and corridors as undirected edges. A state records the agent's current vertex. An action moves to a connected vertex when the corridor is not blocked. The goal test succeeds when the current vertex equals the selected exit.")
    doc.add_heading("2 1 Transition Cost", level=2)
    add_equation(doc, "cost.png", 4.9)
    add_body(doc, "Here d is corridor distance, r is the mean hazard probability at the two endpoints, q is congestion, and the two weights express the selected safety priorities. Every added term is non-negative, so physical distance remains a lower bound for the weighted cost.")
    doc.add_heading("2 2 Environment Classification", level=2)
    add_table(doc, ["Property", "Classification", "Reason"], [
        ["Observability", "Partially observable", "Sensors provide uncertain evidence"],
        ["Change", "Dynamic", "A corridor can close after movement begins"],
        ["Outcome", "Nondeterministic", "A planned action may become unavailable"],
        ["Sequence", "Sequential", "Each move changes the next decision state"],
        ["Agents", "Single and multiple", "The application supports individual and coordinated routes"],
    ], [1.25, 1.55, 3.75])

    doc.add_heading("3 Search Algorithms", level=1)
    doc.add_heading("3 1 Breadth First Search", level=2)
    add_body(doc, "BFS stores the frontier in a first-in, first-out queue and expands states by depth. It is complete on this finite graph. It minimizes the number of edges when actions have equal cost, but it does not minimize the weighted evacuation cost.")
    doc.add_heading("3 2 Depth First Search", level=2)
    add_body(doc, "DFS stores the frontier in a last-in, first-out stack. It follows one branch before backtracking. The result depends on neighbour order and may be longer or less safe than alternatives. The iterative implementation avoids recursion-depth limits.")
    doc.add_heading("3 3 Greedy Best First Search", level=2)
    add_body(doc, "Greedy search uses a priority queue ordered by straight-line distance to the goal. It often expands few nodes, but it ignores accumulated distance, hazard, and congestion. A visually direct route can therefore have a high weighted cost.")
    doc.add_heading("3 4 A Star Search", level=2)
    add_equation(doc, "astar.png", 2.2)
    add_body(doc, "The value g stores the best known accumulated weighted cost and h stores straight-line distance to the exit. Straight-line distance is admissible when map edge distances are consistent with the coordinates because hazard and congestion only increase cost. The implementation updates parent links and g scores whenever it finds a cheaper path.")
    add_table(doc, ["Algorithm", "Frontier", "Decision value", "Main limitation"], [
        ["BFS", "Queue", "Depth", "Ignores weighted cost"],
        ["DFS", "Stack", "Last discovered", "Depends on neighbour order"],
        ["Greedy", "Priority queue", "h", "Ignores accumulated cost"],
        ["A star", "Priority queue", "g plus h", "Depends on model and heuristic"],
    ], [1.0, 1.35, 1.4, 2.75])

    doc.add_heading("4 Uncertain Knowledge and Bayesian Reasoning", level=1)
    add_equation(doc, "bayes.png", 5.6)
    add_body(doc, "The prior probability represents the hazard belief before new evidence. The smoke and heat likelihoods describe how probable each sensor report is when a hazard exists or does not exist. The application applies the two updates sequentially under a conditional-independence assumption. The posterior probability is assigned to the selected node and influences subsequent route costs.")
    add_body(doc, "The conditional-independence assumption is a teaching simplification. Real sensors may share causes, require calibration, and produce correlated errors.")

    doc.add_heading("5 Planning in a Changing Environment", level=1)
    add_body(doc, "The replanning drill executes the first movement of the original A star route, closes the next planned corridor, and searches again from the current location. This observe, update, and plan cycle demonstrates nondeterministic planning. Searching from the original start would be incorrect because the agent has already changed state.")
    image = doc.add_picture(str(SCREENSHOTS / "replanning.png"), width=Inches(2.65))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_caption(doc, "Figure 1  Dynamic blockage and replanning view")

    doc.add_heading("6 Multiagent Planning", level=1)
    add_body(doc, "The centralized planner evaluates every feasible assignment of selected agents to available exits. A capacity constraint limits the number of agents assigned to one exit. The objective adds individual route costs and a penalty for corridors used by more than one agent.")
    add_equation(doc, "multiagent.png", 5.4)
    add_body(doc, "This exhaustive method gives an exact assignment for the small classroom graph. Its assignment count grows exponentially, so a large deployment would require a scalable method such as prioritized planning or conflict-based search.")
    image = doc.add_picture(str(SCREENSHOTS / "multi-agent.png"), width=Inches(2.75))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_caption(doc, "Figure 2  Coordinated planning for three agents")

    doc.add_heading("7 Genetic Resource Placement", level=1)
    add_body(doc, "A chromosome contains the campus nodes selected for emergency resources. Fitness is the average risk-aware A star cost from every demand node to its nearest selected resource. Lower fitness values indicate better access.")
    add_table(doc, ["Operator", "Implementation"], [
        ["Initialization", "Unique random location tuples"],
        ["Selection", "Three-candidate tournament"],
        ["Crossover", "Sample locations from the union of two parents"],
        ["Mutation", "Replace one selected location with an unused candidate"],
        ["Elitism", "Copy the two best chromosomes to the next generation"],
        ["Reproducibility", "Fixed random seed with deterministic population order"],
    ], [1.55, 5.0])
    add_body(doc, "A genetic algorithm does not prove global optimality. The application therefore shows its convergence history and uses a fixed seed so an evaluator can reproduce the same classroom experiment.")
    image = doc.add_picture(str(SCREENSHOTS / "genetic.png"), width=Inches(2.75))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_caption(doc, "Figure 3  Resource placement and convergence view")

    doc.add_heading("8 System Design and Implementation", level=1)
    add_table(doc, ["Module", "Responsibility"], [
        ["models.py", "Graph, edge, and search-result data structures"],
        ["algorithms.py", "BFS, DFS, Greedy, and A star"],
        ["risk.py", "Bayesian probability update"],
        ["planning.py", "Unexpected blockage simulation"],
        ["multiagent.py", "Joint exit assignment and conflict penalty"],
        ["genetic.py", "Emergency resource placement"],
        ["scenarios.py", "Preset maps and editable-table conversion"],
        ["visualization.py", "Plotly campus figures"],
        ["app.py", "Streamlit controls and results"],
    ], [1.75, 4.8])
    add_body(doc, "The search algorithms are implemented directly. Streamlit provides the interface, Pandas provides editable and comparison tables, and Plotly renders the graph. Route calculation does not call a library shortest-path function.")
    doc.add_heading("8 1 Editable Campus Map", level=2)
    add_body(doc, "The node table stores names, coordinates, and hazard probabilities. The corridor table stores endpoints, distance, congestion, and blocked status. Validation rejects blank or duplicate names, unknown endpoints, non-positive distances, and probabilities outside zero to one. The custom graph remains separate from the preset scenarios.")

    doc.add_heading("9 Experimental Evaluation", level=1)
    add_body(doc, "The experiment script runs four algorithms on five scenarios with hazard weight 5 and congestion weight 2. The table below shows all algorithms for normal operation and the Library fire scenario. Runtime is omitted because the graph is small enough for interpreter and measurement noise to dominate.")
    with (ROOT / "outputs" / "experiment_results.csv").open(encoding="utf-8") as handle:
        results = list(csv.DictReader(handle))
    selected = [row for row in results if row["scenario"] in ("Normal operation", "Library fire")]
    add_table(doc, ["Scenario", "Algorithm", "Cost", "Distance", "Risk", "Expanded"], [
        [row["scenario"], row["algorithm"], row["weighted_cost"], row["distance"], row["risk_exposure"], row["nodes_expanded"]]
        for row in selected
    ], [1.55, 1.0, 0.8, 0.8, 0.7, 0.75])
    add_body(doc, "In normal operation, A star produced the lowest weighted cost, 17.5, while all four routes had distance 13. In the Library fire scenario, BFS and Greedy chose the ten-unit route through Library and incurred risk exposure 4.075. A star selected a longer route with zero risk exposure and reduced weighted cost from 33.875 to 32.1. DFS also reduced risk exposure, but its congestion and distance produced a higher total cost of 37.0.")

    doc.add_page_break()
    doc.add_heading("9 1 Multiagent Result", level=2)
    with (ROOT / "work" / "experiment_summary.json").open(encoding="utf-8") as handle:
        summary = json.load(handle)
    add_table(doc, ["Agent", "Start", "Assigned exit", "Cost", "Route"], [
        [item["agent"], item["start"], item["exit"], item["cost"], " to ".join(item["route"])]
        for item in summary["multiagent"]["plans"]
    ], [0.75, 1.15, 1.15, 0.65, 2.85])
    add_body(doc, f"The coordinated case achieved joint objective {summary['multiagent']['objective']} with {summary['multiagent']['shared_corridors']} shared corridors. Two agents used North Gate without sharing a corridor, and the Science Lab agent used South Exit.")
    doc.add_heading("9 2 Genetic Result", level=2)
    genetic = summary["genetic"]
    add_body(doc, f"With two resource stations, a population of four, 35 generations, and fixed seed 5, the best average access cost improved from {genetic['initial_best']} to {genetic['final_best']}. The selected locations were {genetic['locations'][0]} and {genetic['locations'][1]}. The result is reproducible but remains a search result rather than a proof of global optimality.")

    doc.add_heading("10 Testing", level=1)
    add_table(doc, ["Test", "Expected property"], [
        ["All algorithms reach an exit", "Valid path starts and ends at requested nodes"],
        ["Severe fire avoidance", "Risk-aware A star avoids Library"],
        ["Blocked corridor handling", "Returned path does not use blocked edges"],
        ["Bayesian evidence", "Positive evidence increases hazard probability"],
        ["No-path handling", "Failure returns an empty path"],
        ["Dynamic replanning", "Revised route excludes the newly blocked corridor"],
        ["Multiagent capacity", "No exit exceeds its configured load"],
        ["Genetic reproducibility", "Fixed seed repeats the same result"],
        ["Map round trip", "Editable tables rebuild the original graph"],
    ], [2.35, 4.2])
    add_body(doc, "All nine automated tests pass. Browser verification also confirmed the live map, replanning, multiagent, genetic, and map-editor tabs after a fresh server restart.")

    doc.add_heading("11 Limitations and Responsible Use", level=1)
    add_bullets(doc, [
        "The campus graph and sensor likelihoods are simulated rather than measured",
        "Straight-line coordinates simplify real building geometry",
        "The model does not represent mobility needs, smoke propagation, speed, panic, or group behaviour",
        "The exhaustive multiagent assignment supports a small academic graph",
        "Sensor evidence uses a simplified conditional-independence assumption",
        "A real system would require certified maps, calibrated devices, offline operation, security controls, and approval from emergency authorities",
    ])
    add_body(doc, "The interface identifies the project as an educational simulator. Users must follow official emergency procedures rather than relying on this prototype during a real incident.")

    doc.add_heading("12 Conclusion", level=1)
    add_body(doc, "SafeRoute AI implements the planned academic scope in one explainable application. It connects state-space representation with uninformed and informed search, converts uncertain evidence into route cost, revises plans after an unexpected change, coordinates several agents, and applies a genetic algorithm to resource placement. The experiment results show why the shortest route and the lowest-risk weighted route can differ. The code, tests, experiment generator, report, presentation, demonstration script, and viva guide form a reproducible submission package.")

    doc.add_heading("References", level=1)
    references = [
        "Russell, S. J., and Norvig, P. Artificial Intelligence A Modern Approach. Pearson.",
        "Python Software Foundation. Python documentation. https://docs.python.org/3/",
        "Streamlit. Streamlit documentation. https://docs.streamlit.io/",
        "Plotly. Plotly Python documentation. https://plotly.com/python/",
        "Python Software Foundation. heapq priority queue documentation. https://docs.python.org/3/library/heapq.html",
    ]
    for index, reference in enumerate(references, 1):
        add_body(doc, reference, f"{index}. ")

    doc.add_heading("Appendix A Local Execution", level=1)
    add_body(doc, "Create and activate a Python virtual environment, install the declared requirements, then run Streamlit from the project directory.")
    for command in (
        "python3 -m venv .venv",
        "source .venv/bin/activate",
        "python -m pip install -r requirements.txt",
        "streamlit run app.py",
        "pytest -q",
        "python scripts/run_experiments.py",
    ):
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(command)
        run.font.name = "Liberation Mono"
        run.font.size = Pt(9.5)

    OUTPUT.parent.mkdir(exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
